import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from docx import Document
from swarm_world import A, B, Paper, build
from test_swarm import critic, judge, literature, proposal, three_animal_corpus

import chronobio.report as report
from chronobio.export import build_records
from chronobio.figures import graph_svg, hypothesis_matrix_svg, legend_height, scope_ids, workflow_svg
from chronobio.graph import snapshot
from chronobio.layout import layout
from chronobio.llm import ScriptedLLM
from chronobio.models import Publication
from chronobio.render import find_chrome, html_to_pdf, svg_to_png
from chronobio.swarm import load_result, run_swarm

NS = "{http://www.w3.org/2000/svg}"
needs_chrome = pytest.mark.skipif(find_chrome() is None, reason="Chrome not installed")


def accepted(tmp_path):
    corpus = three_animal_corpus(tmp_path)
    llm = ScriptedLLM([proposal(corpus.claim_ids), critic("accept"), judge("yes")])
    return corpus, run_swarm(llm, corpus.store, corpus.settings, literature([]), max_gaps=1)


def unaccepted(tmp_path):
    corpus = three_animal_corpus(tmp_path)
    llm = ScriptedLLM([proposal(corpus.claim_ids[:1]), critic("revise")] * 3)
    return corpus, run_swarm(llm, corpus.store, corpus.settings, literature([]), max_gaps=1)


def texts(root) -> str:
    return " ".join("".join(t.itertext()) for t in root.iter(NS + "text"))


# ── figures ──────────────────────────────────────────────────────────────────

def test_hypothesis_matrix_labels_every_cell_in_words_and_has_no_combined_score(tmp_path):
    corpus, result = unaccepted(tmp_path)
    root = ET.fromstring(hypothesis_matrix_svg(build_records(corpus.store, result)))
    body = texts(root)
    for word in ("proposed", "revise", "not performed", "not searched", "none is a probability"):
        assert word in body
    assert "score:" not in body.lower() and root.get("role") == "img"
    assert {g.get("data-hypothesis") for g in root.iter(NS + "g") if g.get("data-hypothesis")} == {r.id for r in result.hypotheses}


def test_every_hypothesis_is_rated_even_if_not_accepted(tmp_path):
    """Regression: unaccepted hypotheses showed 'not assessed' for dimensions that can be computed."""
    corpus, result = unaccepted(tmp_path)
    (rec,) = build_records(corpus.store, result)
    assert result.cards == [] and rec["ranking"]["evidence_strength"]["level"] != "not_assessed"
    assert rec["ranking"]["feasibility"]["level"] == "not_assessed"


def test_workflow_shows_where_a_hypothesis_stopped_and_what_was_not_done(tmp_path):
    ok_corpus, ok = accepted(tmp_path / "ok")
    done = texts(ET.fromstring(workflow_svg(build_records(ok_corpus.store, ok)[0])))
    assert "accept after 1 round(s)" in done and "known" in done
    assert "NOT PERFORMED" in done and "not designed yet" in done
    bad_corpus, bad = unaccepted(tmp_path / "bad")
    stuck = texts(ET.fromstring(workflow_svg(build_records(bad_corpus.store, bad)[0])))
    assert "revise after 3 round(s)" in stuck and "unresolved" in stuck and "not run" in stuck


@pytest.mark.parametrize("width", [960, 720, 520])
def test_graph_legend_never_runs_off_the_figure(tmp_path, width):
    """Regression: the last legend entry was clipped at the right edge."""
    corpus, _ = accepted(tmp_path)
    from chronobio.graph import load_graph
    nodes_list, edges = load_graph(corpus.store)
    snap = snapshot(nodes_list, edges, corpus.claims, corpus.pubs, None)
    placed = layout([n.id for n in snap.nodes], [(e.source, e.target) for e in snap.edges], width=width)
    height = placed.height + legend_height(width)
    root = ET.fromstring(graph_svg(snap, placed.pos, {n.id: n for n in nodes_list}, title="t", width=width, height=height))
    legend = next(g for g in root.iter(NS + "g") if g.get("class") == "legend")
    box = legend.find(NS + "rect")
    assert float(box.get("y")) >= placed.height - 4 and float(box.get("y")) + float(box.get("height")) <= height
    for t in legend.iter(NS + "text"):
        assert float(t.get("x")) + 6.2 * len("".join(t.itertext())) <= width


def test_scope_ids_makes_inline_figures_safe_to_share_a_page():
    svg = '<svg aria-labelledby="t d"><title id="t">x</title><desc id="d">y</desc><marker id="m"/><path marker-end="url(#m)"/></svg>'
    a, b = scope_ids(svg, "one"), scope_ids(svg, "two")
    ids = re.findall(r'\bid="([^"]+)"', a + b)
    assert len(ids) == len(set(ids)) == 6
    assert 'aria-labelledby="t-one d-one"' in a and "url(#m-two)" in b


# ── report ───────────────────────────────────────────────────────────────────

def build_no_chrome(corpus, tmp_path, monkeypatch, **kw):
    monkeypatch.setattr(report, "find_chrome", lambda: None)
    return report.build_report(corpus.store, corpus.settings, tmp_path / "report", title="Test report", **kw)


def test_report_writes_every_format_and_says_what_it_could_not_make(tmp_path, monkeypatch):
    corpus, _ = accepted(tmp_path)
    res = build_no_chrome(corpus, tmp_path, monkeypatch)
    assert res.html.exists() and res.docx.exists() and res.json.exists() and res.pdf is None
    assert any("no Chrome" in n for n in res.notes)
    assert all((tmp_path / "report" / "figures" / f"{f.key}.svg").exists() for f in res.figures)
    html = res.html.read_text()
    assert "Not produced in this run" in html and "PNG figures and PDF" in html            # stated, not hidden


def test_report_contents_cover_the_spec(tmp_path, monkeypatch):
    corpus, result = accepted(tmp_path)
    res = build_no_chrome(corpus, tmp_path, monkeypatch)
    html = res.html.read_text()
    for section in ("1. Summary", "3. Research gaps", "4. Hypotheses", "5. Competing hypotheses",
                    "6. Validation and experiment design", "7. Limitations", "8. References", "Prior-work audit"):
        assert section in html, section
    (hyp,) = result.hypotheses
    assert hyp.statement in html and "Supporting evidence" in html and "Opposing evidence" in html
    assert "NOT PERFORMED" in html and "Absence of search hits is not evidence of novelty" in html
    assert "[1]" in html and "doi:10.1/0" in html                                          # numbered references with identifiers
    data = json.loads(res.json.read_text())
    assert data["computational_validation"] == "NOT PERFORMED" and data["hypotheses"][0]["statement"] == hyp.statement
    assert data["figures"] and data["references"]


def test_every_figure_lists_claims_that_exist(tmp_path, monkeypatch):
    corpus, _ = accepted(tmp_path)
    res = build_no_chrome(corpus, tmp_path, monkeypatch)
    sources = json.loads((tmp_path / "report" / "figures" / "figure_sources.json").read_text())
    assert set(sources) == {f.key for f in res.figures} and "evidence-graph" in sources
    assert all(cid in corpus.claims for ids in sources.values() for cid in ids)
    html = res.html.read_text()
    assert html.count("<figcaption>") == len(res.figures) and "Claims:" in html


def test_word_report_has_headings_tables_and_a_visible_note_where_figures_are_missing(tmp_path, monkeypatch):
    corpus, _ = accepted(tmp_path)
    res = build_no_chrome(corpus, tmp_path, monkeypatch)
    doc = Document(str(res.docx))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading 1")]
    assert "1. Summary" in headings and "8. References" in headings and len(doc.tables) >= 5
    assert any("Figure not embedded" in p.text for p in doc.paragraphs)                    # not silently dropped


def test_references_without_identifiers_are_flagged_as_unverified(tmp_path, monkeypatch):
    corpus, _ = accepted(tmp_path)
    corpus.store.save(Publication(title="Unidentifiable preprint", year=2020, provenance=build(tmp_path / "x", []).store and None or __import__("swarm_world").PROV))
    res = build_no_chrome(corpus, tmp_path, monkeypatch)
    assert "no DOI or PMID: not verified" in res.html.read_text()
    assert any("no DOI or PMID" in p.text for p in Document(str(res.docx)).paragraphs)


def test_hostile_and_control_characters_cannot_break_html_or_word(tmp_path, monkeypatch):
    corpus, result = accepted(tmp_path)
    hyp = result.hypotheses[0]
    corpus.store.save(hyp.revise(statement=f"{A} <script>alert(1)</script> drives {B}\x07 activity", mechanism="a\x00b"))
    res = build_no_chrome(corpus, tmp_path, monkeypatch)
    html = res.html.read_text()
    assert "<script" not in html and "&lt;script&gt;alert(1)" in html
    assert "script-src" not in html and "default-src 'none'" in html                       # report ships no script at all
    Document(str(res.docx))                                                                # still opens


def test_report_html_has_unique_ids_so_figure_labels_resolve(tmp_path, monkeypatch):
    corpus, _ = accepted(tmp_path)
    html = build_no_chrome(corpus, tmp_path, monkeypatch).html.read_text()
    ids = re.findall(r'\bid="([^"]+)"', html)
    assert len(ids) == len(set(ids))
    for group in re.findall(r'aria-labelledby="([^"]+)"', html):
        assert all(i in ids for i in group.split())


def test_empty_corpus_still_produces_an_honest_report(tmp_path, monkeypatch):
    corpus = build(tmp_path, [])
    res = build_no_chrome(corpus, tmp_path, monkeypatch)
    html = res.html.read_text()
    assert "No hypotheses have been generated" in html and "NOT PERFORMED" in html


def test_unaccepted_hypotheses_stay_in_the_report_with_their_reasons(tmp_path, monkeypatch):
    corpus, result = unaccepted(tmp_path)
    html = build_no_chrome(corpus, tmp_path, monkeypatch).html.read_text()
    assert result.hypotheses[0].statement in html and "proposed" in html and "Critic findings" in html


# ── rendering through Chrome ─────────────────────────────────────────────────

def test_render_commands_use_a_throwaway_profile_and_fail_soft(tmp_path, monkeypatch):
    seen = {}

    def fake_run(cmd, **kwargs):
        seen["cmd"] = cmd
        Path([c for c in cmd if c.startswith("--print-to-pdf=")][0].split("=", 1)[1]).write_bytes(b"%PDF-fake")

    from chronobio import render
    monkeypatch.setattr(render.subprocess, "run", fake_run)
    page = tmp_path / "r.html"
    page.write_text("<p>x</p>")
    assert html_to_pdf(page, tmp_path / "r.pdf", chrome=Path("/bin/echo"))
    cmd = seen["cmd"]
    assert "--headless=new" in cmd and any(c.startswith("--user-data-dir=") for c in cmd) and "--no-pdf-header-footer" in cmd
    assert cmd[-1].startswith("file://")                                                  # local files only


def test_a_missing_browser_binary_means_false_not_a_crash(tmp_path):
    assert svg_to_png("<svg/>", tmp_path / "x.png", 10, 10, chrome=Path("/nonexistent/chrome")) is False
    assert html_to_pdf(tmp_path / "none.html", tmp_path / "x.pdf", chrome=Path("/nonexistent/chrome")) is False


@needs_chrome
def test_real_chrome_rasterizes_svg_at_double_resolution(tmp_path):
    from PIL import Image
    svg = '<svg xmlns="http://www.w3.org/2000/svg" width="120" height="60" viewBox="0 0 120 60"><rect width="120" height="60" fill="#09c"/></svg>'
    assert svg_to_png(svg, tmp_path / "t.png", 120, 60)
    assert Image.open(tmp_path / "t.png").size == (240, 120)


@needs_chrome
def test_full_report_with_chrome_has_pdf_pngs_and_embedded_figures(tmp_path):
    corpus, _ = accepted(tmp_path)
    res = report.build_report(corpus.store, corpus.settings, tmp_path / "report", title="Full")
    assert res.notes == [] and res.pdf and res.pdf.read_bytes().startswith(b"%PDF")
    assert all(f.png and f.png.exists() for f in res.figures)
    assert len(Document(str(res.docx)).inline_shapes) == len(res.figures)


def test_command_line_builds_a_report_and_a_viewer_from_the_store(tmp_path, monkeypatch, capsys):
    from chronobio.__main__ import main
    corpus, _ = accepted(tmp_path)
    monkeypatch.setenv("CHRONOBIO_DATA_DIR", str(corpus.settings.data_dir))
    monkeypatch.setattr(report, "find_chrome", lambda: None)
    assert main(["chronobio", "report", "--out", str(tmp_path / "r"), "--title", "CLI report", "--no-pdf"]) == 0
    out = capsys.readouterr().out
    assert "report.html" in out and "pdf   not produced" in out and (tmp_path / "r" / "report.docx").exists()
    assert main(["chronobio", "viewer", "--out", str(tmp_path / "v")]) == 0
    assert (tmp_path / "v" / "viewer.html").exists()


# ── validation and experiment plan flow through everything ───────────────────

def validated(tmp_path, monkeypatch, effect=2.0):
    from test_validation import COMMON, PARAMS, write_csv
    from chronobio.datasets import ingest_csv
    from chronobio.experiment import design_experiment
    from chronobio.validation import execute_runs, plan_run
    corpus, result = accepted(tmp_path)
    hyp = result.hypotheses[0]
    ds = ingest_csv(corpus.store, write_csv(tmp_path / "synthetic.csv", effect=effect), **COMMON)
    run = execute_runs(corpus.store, [plan_run(corpus.store, hyp, ds, "two_group_numeric", PARAMS).id])[0]
    corpus.store.save(design_experiment(corpus.store.get(hyp.id), corpus.claims, corpus.pubs, [run]))
    monkeypatch.setattr(report, "find_chrome", lambda: None)
    return corpus, hyp, run


def test_a_consistent_run_updates_status_figures_export_and_report(tmp_path, monkeypatch):
    corpus, hyp, run = validated(tmp_path, monkeypatch)
    assert run.verdict == "consistent"
    res = report.build_report(corpus.store, corpus.settings, tmp_path / "r", title="V")
    (rec,) = json.loads(res.json.read_text())["hypotheses"]
    assert rec["status"] == "computationally_tested" and rec["computational_validation"] == "EXECUTED"
    assert rec["validation"][0]["planned_at"] <= rec["validation"][0]["recorded_at"] and rec["experiment_plan"]["power_plan"]
    html = res.html.read_text()
    assert "1 run(s) executed on user datasets" in html and "Pre-registered" in html and "does not validate a hypothesis" in html
    assert "Experiment plan for hypothesis 1" in html and "Computational validation: NOT PERFORMED" not in html
    flow = texts(ET.fromstring(workflow_svg(rec)))
    assert "executed: consistent" in flow and "plan drafted" in flow and "NOT PERFORMED" not in flow
    assert "consistent" in texts(ET.fromstring(hypothesis_matrix_svg([rec])))
    headings = [p.text for p in Document(str(res.docx)).paragraphs if p.style.name.startswith("Heading")]
    assert any(h.startswith("Experiment plan for hypothesis 1") for h in headings)


def test_an_inconclusive_run_leaves_the_status_and_the_report_says_so(tmp_path, monkeypatch):
    corpus, hyp, run = validated(tmp_path, monkeypatch, effect=0.0)
    assert run.verdict == "inconclusive"
    (rec,) = build_records(corpus.store, load_result(corpus.store))
    assert rec["status"] == "evidence_supported" and "executed: inconclusive" in texts(ET.fromstring(workflow_svg(rec)))
