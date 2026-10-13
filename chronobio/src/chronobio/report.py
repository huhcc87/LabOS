"""Research report: one document model, rendered to HTML, Word, PDF and JSON.

Everything in the report is read from stored records; nothing is generated here.
Evidence statements carry a numbered reference, references with no DOI or PMID
are flagged as unverified, every figure lists the claim ids it draws on, and
the limitations section is computed from the data (for example how many claims
come only from abstracts).
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document as WordDocument
from docx.shared import Inches

from .config import Settings
from .export import build_records
from .figures import (
    graph_svg, hypothesis_matrix_svg, legend_height, matrix_svg, scope_ids, timeline_svg, workflow_svg,
)
from .graph import load_graph, restrict, snapshot, suggest_inferred_edges
from .layout import layout
from .models import EvidenceClaim, Publication, RetractionStatus
from .render import find_chrome, html_to_pdf, svg_to_png
from .store import Store
from .swarm import load_result

REPORT_VERSION = "report-v1"


@dataclass
class Figure:
    key: str
    title: str
    svg: str
    width: int
    height: int
    caption: str
    legend: str
    claim_ids: list[str] = field(default_factory=list)
    png: Path | None = None


# A block is a tuple whose first item names its kind; both renderers understand the same few kinds.
Block = tuple


@dataclass
class ReportResult:
    html: Path
    docx: Path
    json: Path
    pdf: Path | None
    figures: list[Figure]
    notes: list[str]                   # anything that could not be produced, stated plainly


def _author_list(pub: Publication) -> str:
    names = pub.authors[:6]
    return ", ".join(names) + (" et al." if len(pub.authors) > 6 else "")


def _reference(pub: Publication) -> str:
    ids = " ".join(x for x in (f"doi:{pub.doi}" if pub.doi else "", f"PMID {pub.pmid}" if pub.pmid else "") if x)
    parts = [p for p in (_author_list(pub), pub.title, pub.journal, str(pub.year or "n.d."), ids) if p]
    return ". ".join(parts) + "."


def _svg_size(svg: str) -> tuple[int, int]:
    m = re.search(r'viewBox="0 0 (\d+) (\d+)"', svg)
    return (int(m.group(1)), int(m.group(2))) if m else (960, 600)


def build_figures(
    store: Store, records: list[dict], nodes_list, edges, claims: dict[str, EvidenceClaim], pubs: dict[str, Publication],
) -> list[Figure]:
    nodes = {n.id: n for n in nodes_list}
    snap = snapshot(nodes_list, edges + suggest_inferred_edges(edges), claims, pubs, None)
    placed = layout([n.id for n in snap.nodes], [(e.source, e.target) for e in snap.edges])
    w, h = placed.width, placed.height + legend_height(placed.width)
    figures: list[Figure] = []

    def add(key, title, svg, caption, legend, claim_ids=()):
        width, height = _svg_size(svg)
        figures.append(Figure(key, title, svg, width, height, caption, legend, sorted(set(claim_ids))))

    graph_legend = ("Solid links cite evidence: green activates, vermilion inhibits, grey is associated or no effect. "
                    "Purple dash-dot links have conflicting evidence. Blue dashed links are hypotheses, not findings.")
    add("evidence-graph", "Evidence graph", graph_svg(snap, placed.pos, nodes, title="Evidence graph", width=w, height=h),
        "Entities and relations reported in the literature corpus, all years.", graph_legend, snap.claim_ids)
    add("timeline", "Evidence by year", timeline_svg(list(claims.values()), pubs),
        "Number of evidence claims per publication year.",
        "Green: supports. Grey: inconclusive. Hatched: contradicts. Claims without a publication year are not shown.",
        claims)
    add("evidence-matrix", "Evidence for and against", matrix_svg(snap, nodes, claims, pubs),
        "Which papers support (S), contradict (C) or are inconclusive about (?) each link.",
        "Columns are papers sorted by year; rows are links.", snap.claim_ids)
    if records:
        add("hypothesis-comparison", "Hypothesis comparison", hypothesis_matrix_svg(records),
            "Hypotheses side by side on separately rated dimensions.",
            "Each cell shows its rating as text. There is no combined score.",
            [c for r in records for c in r["supporting_ids"] + r["opposing_ids"]])
        top = records[0]
        add("hypothesis-workflow", "Hypothesis to experiment", workflow_svg(top),
            "How far the top-listed hypothesis got: critic, prior-work search, validation, experiment design.",
            "Dashed boxes were not done. No validation has been performed.", top["supporting_ids"] + top["opposing_ids"])
        gap = top.get("gap")
        if gap and gap.get("node_ids"):
            sub = restrict(snap, gap["node_ids"])
            sub_pos = layout([n.id for n in sub.nodes], [(e.source, e.target) for e in sub.edges], width=720)
            add("hypothesis-mechanism", "Mechanism around the top hypothesis",
                graph_svg(sub, sub_pos.pos, nodes, title="Mechanism diagram", width=sub_pos.width, height=sub_pos.height + legend_height(sub_pos.width)),
                "The entities and links behind the top-listed hypothesis. Reported links are solid; any hypothesized link is dashed.",
                graph_legend, sub.claim_ids)
    return figures


def _blocks(
    title: str, records: list[dict], figures: list[Figure], claims: dict[str, EvidenceClaim],
    pubs: dict[str, Publication], result, search_logs: int, notes: list[str],
) -> list[Block]:
    fig = {f.key: f for f in figures}
    ref_no: dict[str, int] = {}

    def ref(pub_id: str) -> str:
        if pub_id not in ref_no:
            ref_no[pub_id] = len(ref_no) + 1
        return f"[{ref_no[pub_id]}]"

    def claim_row(cid: str) -> list[str]:
        c = claims[cid]
        pub = pubs.get(c.publication_id)
        return [cid, c.passage, str(pub.year if pub and pub.year else "n.d."), c.design.value,
                str(c.n) if c.n is not None else "", c.effect, ref(c.publication_id)]

    statuses: dict[str, int] = {}
    for r in records:
        statuses[r["status"]] = statuses.get(r["status"], 0) + 1
    abstract_only = sum(c.locator.value == "text_only" for c in claims.values())
    retracted = [p for p in pubs.values() if p.retraction_status == RetractionStatus.RETRACTED]
    years = sorted(p.year for p in pubs.values() if p.year)

    b: list[Block] = [("title", title),
                      ("p", f"Generated {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC from stored evidence "
                            f"({REPORT_VERSION}). Nothing in this report has been validated computationally or experimentally."),
                      ("h1", "1. Summary"),
                      ("bullets", [
                          f"Corpus: {len(pubs)} publication(s)" + (f", {years[0]}–{years[-1]}" if years else "") + f"; {len(claims)} extracted claim(s).",
                          f"Research gaps found: {len(result.gaps)}.",
                          f"Hypotheses: {len(records)} (" + (", ".join(f"{n} {s.replace('_', ' ')}" for s, n in sorted(statuses.items())) or "none") + ").",
                          ("Computational validation: " + (f"{sum(1 for r in records for v in r['validation'] if v['status'] == 'executed')} run(s) executed on user datasets (see section 6)" if any(v['status'] == 'executed' for r in records for v in r['validation']) else "NOT PERFORMED") + ". Experimental validation: NOT PERFORMED."),
                      ])]
    if notes:
        b += [("p", "Not produced in this run: " + "; ".join(notes))]

    b += [("h1", "2. Evidence landscape")]
    for key in ("evidence-graph", "timeline", "evidence-matrix"):
        if key in fig:
            b.append(("figure", fig[key]))

    b += [("h1", "3. Research gaps"),
          ("table", ["Kind", "Description", "Papers", "Priority"],
           [[g.kind.value, g.description, str(g.paper_count), f"{g.priority:.2f}"] for g in result.gaps] or [["none", "", "", ""]])]

    b += [("h1", "4. Hypotheses")]
    if not records:
        b.append(("p", "No hypotheses have been generated for this corpus."))
    for i, r in enumerate(records, start=1):
        rank = r.get("ranking") or {}
        verdict = (r.get("critic") or {}).get("verdict", "not reviewed")
        b += [("h2", f"4.{i} {r['statement']}"),
              ("p", f"Status: {r['status'].replace('_', ' ')} · Critic: {verdict} after {len(r['verdict_history'])} round(s) · "
                    f"Prior-work search: {r['novelty_status'].replace('_', ' ')} · Validation: {r['computational_validation']}")]
        if r.get("gap"):
            b.append(("p", f"Gap ({r['gap']['kind']}): {r['gap']['description']}"))
        if r["mechanism"]:
            b.append(("p", f"Proposed mechanism: {r['mechanism']}"))
        b += [("h3", "Predictions"), ("bullets", r["predictions"]), ("h3", "Would be falsified by"), ("bullets", r["falsifiers"])]
        if r["alternative_mechanisms"]:
            b += [("h3", "Alternative explanations"), ("bullets", r["alternative_mechanisms"])]
        b += [("h3", "Supporting evidence"),
              ("table", ["Claim", "Quoted passage", "Year", "Design", "n", "Effect", "Source"],
               [claim_row(c) for c in r["supporting_ids"] if c in claims]),
              ("h3", "Opposing evidence")]
        b.append(("table", ["Claim", "Quoted passage", "Year", "Design", "n", "Effect", "Source"],
                  [claim_row(c) for c in r["opposing_ids"] if c in claims]) if r["opposing_ids"]
                 else ("p", "None found in this corpus."))
        b += [("h3", "Uncertainty"), ("bullets", r["uncertainty"] or ["Not stated."])]
        if r["critic"]:
            findings = r["critic"]["rule_findings"] + r["critic"]["notes"]
            b += [("h3", "Critic findings"), ("bullets", findings or ["None."])]
        if r["novelty"]:
            n = r["novelty"]
            b += [("h3", f"Prior-work audit ({n['status'].replace('_', ' ')})"),
                  ("table", ["Query", "Run on", "Hits", "Examined"],
                   [[s["query"], s["run_at"][:10], str(s["hits"]), str(s["examined"])] for s in n["searches"]]),
                  ("bullets", [f"Prior work: {q}" for q in n["prior_work_quotes"]] + [f"Limitation: {x}" for x in n["limitations"]])]
        if rank:
            b += [("h3", "Ranking (rated separately, no combined score)"),
                  ("table", ["Dimension", "Rating", "Reason"],
                   [[k.replace("_", " "), v["level"].replace("_", " "), "; ".join(v["reasons"])] for k, v in rank.items()])]
    if records:
        b.append(("figure", fig["hypothesis-comparison"]))
        for key in ("hypothesis-workflow", "hypothesis-mechanism"):
            if key in fig:
                b.append(("figure", fig[key]))

    by_gap: dict[str, list[dict]] = {}
    for r in records:
        if r.get("gap"):
            by_gap.setdefault(r["gap"]["id"], []).append(r)
    competing = [(gid, rs) for gid, rs in by_gap.items() if len(rs) > 1]
    b.append(("h1", "5. Competing hypotheses"))
    b.append(("table", ["Gap", "Hypotheses addressing it"],
              [[rs[0]["gap"]["description"], " | ".join(r["statement"] for r in rs)] for _, rs in competing])
             if competing else ("p", "Each gap has one hypothesis; none compete."))

    runs = [(r, v) for r in records for v in r["validation"]]
    executed = [v for _, v in runs if v["status"] == "executed"]
    b += [("h1", "6. Validation and experiment design")]
    if not runs:
        b += [("p", "No dataset has been attached and no analysis has been run: computational validation is NOT PERFORMED. "
                    "Each prediction below names a measurement that would test its hypothesis."),
              ("table", ["Hypothesis", "Measurements that would test it"],
               [[r["statement"], "; ".join(r["predictions"])] for r in records] or [["none", ""]])]
    else:
        b += [("p", f"{len(executed)} analysis run(s) executed on the user's own dataset(s). Each endpoint was pre-registered before its "
                    "result existed (see the planned and recorded times). A result is evidence from that dataset only and does not "
                    "validate a hypothesis."),
              ("table", ["Hypothesis", "Dataset", "Method", "Status", "Verdict", "Effect (Hedges g, held-out)", "Adjusted p", "Pre-registered", "Recorded"],
               [[r["statement"][:60], v["dataset"], v["method"], v["status"], v["verdict"] or "-",
                 (f"{v['metrics']['hedges_g_heldout']:.2f} ({v['intervals']['hedges_g_heldout'][0]:.2f} to {v['intervals']['hedges_g_heldout'][1]:.2f})"
                  if v["metrics"] else "-"),
                 f"{v['metrics']['p_adjusted']:.4f}" if v["metrics"] else "-", v["planned_at"][:19], v["recorded_at"][:19]]
                for r, v in runs]),
              ("bullets", [f"{v['dataset']}: {c}" for _, v in runs for c in v["caveats"]])]
    for i, r in enumerate(records, start=1):
        e = r.get("experiment_plan")
        if not e:
            continue
        b += [("h2", f"Experiment plan for hypothesis {i} (draft for expert review)"), ("h3", "Aims"), ("bullets", e["aims"]),
              ("h3", "Models and assays used in the cited studies"), ("bullets", e["models_and_assays"] or ["None stated in the cited evidence."]),
              ("h3", "Controls"), ("bullets", e["controls"]), ("h3", "Endpoints and expected outcomes"), ("bullets", e["expected_outcomes"]),
              ("h3", "Alternative outcomes"), ("bullets", e["alternative_outcomes"] or ["None listed."]),
              ("h3", "Sample size and replication"), ("p", e["power_plan"]), ("p", e["replication"]),
              ("h3", "Go / no-go criteria"), ("bullets", e["go_no_go"]),
              ("h3", "Assumptions"), ("bullets", e["assumptions"] or ["None."]), ("h3", "Ethics, cost and feasibility flags"), ("bullets", e["flags"])]

    limits = [
        f"{abstract_only} of {len(claims)} claims come from abstracts, not full text, so no page-level source highlight exists for them.",
        "Claims and graph links were extracted by a local language model and verified against the source text; entity names are not "
        "matched to ontologies, so near-duplicate names are not merged.",
        f"{len(retracted)} publication(s) are flagged as retracted." if retracted else "No retracted publications were found in this corpus.",
        "Prior-work searches cover Europe PMC titles and abstracts only; full texts and patents were not searched, and absence of hits is not evidence of novelty.",
        f"{search_logs} literature search(es) are recorded; each lists its query and date.",
    ] + notes_for_corpus(pubs)
    b += [("h1", "7. Limitations"), ("bullets", limits)]

    cited = list(ref_no)
    for p in pubs.values():                      # references for every paper, cited or not, so the corpus is auditable
        if p.id not in ref_no:
            cited.append(p.id)
            ref(p.id)
    b += [("h1", "8. References"),
          ("refs", [(ref_no[pid], _reference(pubs[pid]), not (pubs[pid].doi or pubs[pid].pmid)) for pid in cited if pid in pubs])]

    used = {c for r in records for c in r["supporting_ids"] + r["opposing_ids"]}
    b += [("h1", "Appendix A. Evidence behind the hypotheses"),
          ("table", ["Claim", "Quoted passage", "Year", "Design", "n", "Effect", "Source"], [claim_row(c) for c in sorted(used) if c in claims])]
    b += [("h1", "Appendix B. Figure sources"),
          ("table", ["Figure", "Claims drawn on"],
           [[f.title, ", ".join(f.claim_ids[:10]) + (f" … ({len(f.claim_ids)} total)" if len(f.claim_ids) > 10 else "")] for f in figures])]
    return b


def notes_for_corpus(pubs: dict[str, Publication]) -> list[str]:
    missing = sum(1 for p in pubs.values() if not (p.doi or p.pmid))
    return [f"{missing} reference(s) have no DOI or PMID and could not be checked against a registry."] if missing else []


# ── renderers ────────────────────────────────────────────────────────────────

_CSS = """body{font:14px/1.5 system-ui,-apple-system,'Segoe UI',sans-serif;max-width:980px;margin:24px auto;padding:0 20px;color:#1f2937}
h1{font-size:20px;margin:28px 0 8px;border-bottom:1px solid #d1d5db;padding-bottom:4px}h2{font-size:16px;margin:22px 0 6px}h3{font-size:13px;margin:12px 0 4px;color:#374151}
.title{font-size:26px;font-weight:700;margin:0 0 6px}table{border-collapse:collapse;width:100%;font-size:12px;margin:6px 0 12px}
th,td{border:1px solid #d1d5db;padding:4px 6px;text-align:left;vertical-align:top}th{background:#f3f4f6}
figure{margin:14px 0;page-break-inside:avoid}figure svg{max-width:100%;height:auto}figcaption{font-size:12px;color:#4b5563}
.warn{color:#b91c1c;font-weight:600}.small{font-size:12px;color:#6b7280}li{margin:2px 0}
@media print{figure svg{width:auto;height:auto;max-width:100%;max-height:225mm}figure{break-inside:avoid}h1,h2,h3{break-after:avoid}}"""


_BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(text: str) -> str:
    """Word files are XML: control characters (which PDFs sometimes yield) would make them unreadable."""
    return _BAD_XML.sub("", text)


def render_html(title: str, blocks: list[Block]) -> str:
    out: list[str] = []
    for blk in blocks:
        kind = blk[0]
        if kind == "title":
            out.append(f'<p class="title">{escape(blk[1])}</p>')
        elif kind in ("h1", "h2", "h3"):
            out.append(f"<{kind}>{escape(blk[1])}</{kind}>")
        elif kind == "p":
            out.append(f"<p>{escape(blk[1])}</p>")
        elif kind == "bullets":
            out.append("<ul>" + "".join(f"<li>{escape(x)}</li>" for x in blk[1]) + "</ul>")
        elif kind == "table":
            head = "".join(f"<th>{escape(h)}</th>" for h in blk[1])
            rows = "".join("<tr>" + "".join(f"<td>{escape(c)}</td>" for c in row) + "</tr>" for row in blk[2])
            out.append(f"<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>")
        elif kind == "figure":
            f: Figure = blk[1]
            ids = ", ".join(f.claim_ids[:8]) + (f" … ({len(f.claim_ids)} claims)" if len(f.claim_ids) > 8 else "")
            out.append(f"<figure>{scope_ids(f.svg, f.key + str(len(out)))}<figcaption><b>{escape(f.title)}.</b> {escape(f.caption)} {escape(f.legend)} "
                       f'<span class="small">Claims: {escape(ids) or "none"}.</span></figcaption></figure>')
        elif kind == "refs":
            out.append("<ol>" + "".join(
                f'<li value="{n}">{escape(text)}' + (' <span class="warn">⚠ no DOI or PMID: not verified</span>' if bad else "") + "</li>"
                for n, text, bad in blk[1]) + "</ol>")
    body = "\n".join(out)
    csp = "default-src 'none'; style-src 'unsafe-inline'; img-src data:; base-uri 'none'; form-action 'none'"
    return (f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta http-equiv="Content-Security-Policy" content="{csp}">'
            f'<meta name="viewport" content="width=device-width, initial-scale=1"><title>{escape(title)}</title><style>{_CSS}</style></head>'
            f"<body>{body}</body></html>")


def render_docx(blocks: list[Block], out: Path) -> None:
    doc = WordDocument()
    for blk in blocks:
        kind = blk[0]
        if kind == "title":
            doc.add_heading(_clean(blk[1]), level=0)
        elif kind in ("h1", "h2", "h3"):
            doc.add_heading(_clean(blk[1]), level=int(kind[1]))
        elif kind == "p":
            doc.add_paragraph(_clean(blk[1]))
        elif kind == "bullets":
            for item in blk[1]:
                doc.add_paragraph(_clean(item), style="List Bullet")
        elif kind == "table":
            table = doc.add_table(rows=1, cols=len(blk[1]))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, blk[1]):
                cell.paragraphs[0].add_run(_clean(text)).bold = True
            for row in blk[2]:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = _clean(text)
            doc.add_paragraph()
        elif kind == "figure":
            f: Figure = blk[1]
            if f.png and f.png.exists():
                # Fit the page: at most 6.4in wide and 8.5in tall, keeping the aspect ratio.
                doc.add_picture(str(f.png), width=Inches(min(6.4, 8.5 * f.width / f.height)))
            else:
                doc.add_paragraph(f"[Figure not embedded: no PNG renderer available. Editable SVG: figures/{f.key}.svg]")
            ids = ", ".join(f.claim_ids[:8]) + (f" … ({len(f.claim_ids)} claims)" if len(f.claim_ids) > 8 else "")
            doc.add_paragraph().add_run(_clean(f"{f.title}. {f.caption} {f.legend} Claims: {ids or 'none'}.")).italic = True
        elif kind == "refs":
            for n, text, bad in blk[1]:
                doc.add_paragraph(_clean(f"[{n}] {text}" + (" ⚠ no DOI or PMID: not verified" if bad else "")))
    doc.save(str(out))


# ── entry point ──────────────────────────────────────────────────────────────

def build_report(
    store: Store, settings: Settings, out_dir: Path | None = None, *, title: str = "ChronoBio research report", pdf: bool = True,
) -> ReportResult:
    out_dir = out_dir or settings.data_dir / "report"
    (out_dir / "figures").mkdir(parents=True, exist_ok=True)
    pubs = {p.id: p for p in store.list(Publication)}
    claims = {c.id: c for c in store.list(EvidenceClaim)}
    nodes_list, edges = load_graph(store)
    result = load_result(store)
    records = build_records(store, result)
    figures = build_figures(store, records, nodes_list, edges, claims, pubs)

    notes: list[str] = []
    chrome = find_chrome()
    for f in figures:
        (out_dir / "figures" / f"{f.key}.svg").write_text(f.svg)
        if chrome:
            png = out_dir / "figures" / f"{f.key}.png"
            f.png = png if svg_to_png(f.svg, png, f.width, f.height, chrome=chrome) else None
    if not chrome:
        notes.append("PNG figures and PDF (no Chrome found; SVG figures, HTML and Word files were written)")
    elif any(f.png is None for f in figures):
        notes.append("some PNG figures (render failed)")

    from .models import SearchLog
    blocks = _blocks(title, records, figures, claims, pubs, result, len(store.list(SearchLog)), notes)
    html_path = out_dir / "report.html"
    html_path.write_text(render_html(title, blocks))
    docx_path = out_dir / "report.docx"
    render_docx(blocks, docx_path)
    (out_dir / "figures" / "figure_sources.json").write_text(json.dumps({f.key: f.claim_ids for f in figures}, indent=2))
    json_path = out_dir / "report.json"
    json_path.write_text(json.dumps({
        "title": title, "version": REPORT_VERSION, "hypotheses": records,
        "gaps": [{"kind": g.kind.value, "description": g.description, "papers": g.paper_count} for g in result.gaps],
        "references": [_reference(p) for p in pubs.values()],
        "figures": [{"key": f.key, "title": f.title, "caption": f.caption, "legend": f.legend, "claims": f.claim_ids} for f in figures],
        "computational_validation": "EXECUTED" if any(r["computational_validation"] == "EXECUTED" for r in records) else "NOT PERFORMED",
        "experimental_validation": "NOT PERFORMED",
    }, indent=2, ensure_ascii=False))

    pdf_path: Path | None = None
    if pdf and chrome:
        pdf_path = out_dir / "report.pdf"
        if not html_to_pdf(html_path, pdf_path, chrome=chrome):
            pdf_path = None
            notes.append("PDF (render failed)")
    return ReportResult(html_path, docx_path, json_path, pdf_path, figures, notes)
